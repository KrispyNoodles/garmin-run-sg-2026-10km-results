"""
Download the Garmin Run Singapore 21KM (shortened to 10km) results from the
Sporthive/Speedhive results service and save them to an Excel file.

Setup (once):   pip install requests openpyxl
Run:            python sporthive_10km_v2.py
Output:         garmin_run_10km_results.xlsx
                first_record_debug.json   (one raw record, for checking field names)
"""
import json
import sys
import time

import requests
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

RACE_ID = "520d9ead-cf51-4dc9-a563-7fcced16f81c"
URL = f"https://eventresults-api.speedhive.com/sporthive/races/{RACE_ID}/participants"
BASE_PARAMS = {"size": 10, "category": "ALL_RESULTS"}
OUT_FILE = "garmin_run_10km_results.xlsx"
HEADERS = {"User-Agent": "Mozilla/5.0", "Accept": "application/json"}


def get(params):
    r = requests.get(URL, params=params, headers=HEADERS, timeout=30)
    r.raise_for_status()
    return r.json()


def find_items(data):
    """Find the list of participant records in whatever shape the API returns."""
    if isinstance(data, list):
        return data
    if isinstance(data, dict):
        for key in ("participants", "results", "items", "data", "classifications", "content"):
            if isinstance(data.get(key), list):
                return data[key]
        for v in data.values():
            if isinstance(v, list) and v and isinstance(v[0], dict):
                return v
    return []


def flat(d, prefix=""):
    """Flatten nested dicts so keys like 'result.chipTime' can be matched too."""
    out = {}
    for k, v in d.items():
        if isinstance(v, dict):
            out.update(flat(v, f"{prefix}{k}."))
        else:
            out[f"{prefix}{k}"] = v
    return out


def pick(f, *names):
    """Return the first non-empty value whose key (or last key part) matches a name."""
    lower = {k.lower(): v for k, v in f.items()}
    for n in names:
        v = lower.get(n.lower())
        if v not in (None, ""):
            return v
    for n in names:
        for k, v in lower.items():
            if k.split(".")[-1] == n.lower() and v not in (None, ""):
                return v
    return ""


def record_key(item):
    f = flat(item)
    return str(pick(f, "bib", "bibNumber", "id")) + "|" + str(pick(f, "name", "fullName"))


def split_5k(item):
    """Time at the 5 km mark, from legs -> participantSplits."""
    for leg in item.get("legs") or []:
        for s in leg.get("participantSplits") or []:
            if s.get("distanceInMeters") == 5000 or "5k" in str(s.get("splitName", "")).lower():
                return s.get("totalSplitDuration") or s.get("splitDuration") or ""
    return ""


def to_row(item):
    f = flat(item)
    name = pick(f, "name", "fullName", "displayName")
    if not name:
        first = pick(f, "firstName")
        last = pick(f, "lastName", "surname")
        name = f"{first} {last}".strip()
    return [
        name,
        pick(f, "bib", "bibNumber", "startNumber"),
        pick(f, "gender", "sex"),
        pick(f, "category", "categoryName", "raceCategory"),
        pick(f, "gunTimeOfParticipant", "gunTime", "grossTime", "gun"),
        pick(f, "chipTimeOfParticipant", "chipTime", "netTime", "chip", "time"),
        split_5k(item),
    ]


def main():
    first = get(BASE_PARAMS)
    if isinstance(first, dict):
        print("Response keys:", list(first.keys()))
    items = find_items(first)
    if not items:
        print("No participant records found in the response.")
        sys.exit(1)

    with open("first_record_debug.json", "w", encoding="utf-8") as fh:
        json.dump(items[0], fh, indent=2, ensure_ascii=False)

    seen = {record_key(i) for i in items}
    all_items = list(items)
    size = BASE_PARAMS["size"]
    print(f"Page 1: {len(items)} records")

    # Work out how this API pages: try common parameter styles on the 2nd page.
    candidates = [
        {"page": 1}, {"page": 2}, {"offset": size}, {"skip": size}, {"start": size},
    ]
    paging = None
    if len(items) >= size:
        for cand in candidates:
            try:
                data = get({**BASE_PARAMS, **cand})
            except requests.HTTPError:
                continue
            new = [i for i in find_items(data) if record_key(i) not in seen]
            if new:
                paging = cand
                for i in new:
                    seen.add(record_key(i))
                all_items.extend(new)
                print(f"Paging style found: {list(cand)[0]}  (page 2: {len(new)} new)")
                break
        if paging is None:
            print("Could not find a paging parameter; only the first page was saved.")

    if paging:
        key = list(paging)[0]
        step = 1 if key == "page" else size
        value = paging[key]
        while True:
            value += step
            time.sleep(0.3)
            data = get({**BASE_PARAMS, key: value})
            new = [i for i in find_items(data) if record_key(i) not in seen]
            if not new:
                break
            for i in new:
                seen.add(record_key(i))
            all_items.extend(new)
            print(f"Fetched {len(all_items)} participants...")

    rows = [to_row(i) for i in all_items]

    wb = Workbook()
    ws = wb.active
    ws.title = "10km Results"
    header = ["Name", "Bib Number", "Gender", "Category", "Gun Time", "Chip Time", "5K Split Time"]
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    for n, h in enumerate(header, start=1):
        w = max(len(h), *(len(str(r[n - 1])) for r in rows)) + 2
        ws.column_dimensions[get_column_letter(n)].width = min(w, 40)
    ws.freeze_panes = "A2"
    wb.save(OUT_FILE)

    blanks = sum(1 for r in rows if not all(r))
    print(f"Saved {len(rows)} rows to {OUT_FILE}")
    if blanks:
        print(f"{blanks} rows have a blank field. Open first_record_debug.json and "
              "send it to Claude if any column looks wrong.")


if __name__ == "__main__":
    main()
